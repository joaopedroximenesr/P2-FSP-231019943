"""
08_planilha_formulas.py — Gera a planilha 'planilha_calculos_P2_231019943.xlsx' com FÓRMULAS VIVAS do Excel,
reproduzindo deflação, população harmonizada, indicadores, valores reais e per capita de Manaus e Parintins
a partir dos valores nominais extraídos do Siconfi. Permite conferir os cálculos sem Python.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TD = os.path.join(ROOT, "03_dados_tratados")
f = pd.read_csv(os.path.join(TD, "painel_focos_manaus_parintins.csv"))
defl = pd.read_csv(os.path.join(TD, "deflatores.csv"))
pop = pd.read_csv(os.path.join(TD, "populacao_manaus_parintins.csv"))
ibge = pd.read_csv(os.path.join(ROOT, "02_dados_brutos_atualizados/planilhas/ibge_sidra.csv"), dtype=str)

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F5A96"); IN = PatternFill("solid", fgColor="FFF2CC")
def cab(ws, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c); cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.row_dimensions[1].height = 45

# ---- LEIA-ME
ws = wb.active; ws.title = "LEIA-ME"
for i, t in enumerate([
    "Planilha de cálculos — P2 FSP — Matrícula 231019943",
    "Células amarelas = dados de entrada copiados das fontes (Siconfi/DCA, IBGE/Sidra). Demais células = fórmulas.",
    "Aba IPCA: número-índice mensal (Sidra 1737) -> média anual -> fator para R$ de 2025.",
    "Aba Populacao: Censos 2010/2022 e estimativas -> população harmonizada (interpolação geométrica).",
    "Abas Manaus e Parintins: valores nominais (R$) -> receita líquida, transferências líquidas, dependência, autonomia, resultado, valores reais e per capita.",
    "Os mesmos cálculos são feitos em Python (06_codigos_e_planilhas/02_tratamento.py); a aba Conferencia compara os dois resultados.",
], 1):
    ws.cell(row=i, column=1, value=t)
ws.column_dimensions["A"].width = 140

# ---- IPCA mensal e fatores
ws = wb.create_sheet("IPCA")
cab(ws, ["Mês (AAAAMM)", "IPCA número-índice (dez/1993=100)", "", "Ano", "Média anual do índice", "Fator p/ R$ de 2025"])
ip = ibge[(ibge.consulta == "ipca_indice_mensal") & (ibge.periodo >= "201501")].sort_values("periodo")
for i, (_, r) in enumerate(ip.iterrows(), 2):
    ws.cell(row=i, column=1, value=int(r.periodo)); c = ws.cell(row=i, column=2, value=float(r.valor_publicado)); c.fill = IN
n = len(ip) + 1
anos = list(range(2015, 2026))
for k, a in enumerate(anos, 2):
    ws.cell(row=k, column=4, value=a)
    ws.cell(row=k, column=5, value=f'=AVERAGEIFS($B$2:$B${n},$A$2:$A${n},">="&D{k}*100+1,$A$2:$A${n},"<="&D{k}*100+12)')
linha2025 = 2 + anos.index(2025)
for k in range(2, 2 + len(anos)):
    ws.cell(row=k, column=6, value=f"=$E${linha2025}/E{k}")

# ---- População
ws = wb.create_sheet("Populacao")
cab(ws, ["Município", "Ano", "Censo 2010 (1/8/2010)", "Censo 2022 (1/8/2022)", "Estimativa 2024 (1/7/2024)", "Estimativa oficial do ano", "População harmonizada", "Regra"])
row = 2
for mun, cod in [("Manaus", 1302603), ("Parintins", 1303403)]:
    p = pop[pop.cod_ibge == cod].set_index("ano")
    c10 = ibge[(ibge.consulta == "pop_censo2010_1378") & (ibge.cod_ibge == str(cod))].valor_publicado.astype(float).iloc[0]
    c22 = p.loc[2022, "pop_harmonizada"]; e24 = p.loc[2024, "pop_oficial"]
    for a in range(2016, 2026):
        ws.cell(row=row, column=1, value=mun); ws.cell(row=row, column=2, value=a)
        for col, v in [(3, c10), (4, c22), (5, e24)]:
            ws.cell(row=row, column=col, value=float(v)).fill = IN
        of = p.loc[a, "pop_oficial"]
        if pd.notna(of):
            ws.cell(row=row, column=6, value=float(of)).fill = IN
        if a <= 2021:
            ws.cell(row=row, column=7, value=f"=ROUND(C{row}*(D{row}/C{row})^((B{row}+6/12-(2010+7/12))/12),0)"); regra = "Interpolação Censo 2010–2022"
        elif a == 2022:
            ws.cell(row=row, column=7, value=f"=D{row}"); regra = "Censo 2022"
        elif a == 2023:
            ws.cell(row=row, column=7, value=f"=ROUND(D{row}*((E{row}/D{row})^(1/(2024+6/12-(2022+7/12))))^(2023+6/12-(2022+7/12)),0)"); regra = "Interpolação Censo 2022–estimativa 2024"
        else:
            ws.cell(row=row, column=7, value=f"=F{row}"); regra = "Estimativa IBGE pós-censo"
        ws.cell(row=row, column=8, value=regra)
        row += 1

# ---- Municípios
ENT = [("rec_corrente", "Receita corrente bruta (1.0)"), ("ded_rec_corrente", "Deduções da receita corrente (FUNDEB e outras)"),
       ("rec_capital", "Receita de capital (2.0)"), ("rec_tributaria", "Receita tributária (1.1)"),
       ("transf_correntes", "Transferências correntes brutas (1.7)"), ("ded_transf_correntes", "Deduções sobre transf. correntes"),
       ("transf_capital", "Transferências de capital (2.4)"), ("fpm", "FPM"), ("icms_cota", "Cota-parte ICMS"),
       ("fundeb_receb", "FUNDEB recebido"), ("sus_corrente", "SUS corrente"), ("desp_total", "Despesa empenhada exceto intra")]
CALC = [("Receita total líquida", "=B{r}-C{r}+D{r}"), ("Transferências líquidas", "=F{r}-G{r}+H{r}"),
        ("Dependência (%)", "=100*O{r}/N{r}"), ("Autonomia tributária (%)", "=100*E{r}/(B{r}-C{r})"),
        ("Resultado orçamentário (R$)", "=N{r}-M{r}"), ("Fator IPCA (R$ 2025)", "=INDEX(IPCA!$F$2:$F$12,MATCH(A{r},IPCA!$D$2:$D$12,0))"),
        ("População harmonizada", "=SUMIFS(Populacao!$G$2:$G$21,Populacao!$A$2:$A$21,\"{mun}\",Populacao!$B$2:$B$21,A{r})"),
        ("Receita total líquida real pc", "=N{r}*S{r}/T{r}"), ("Receita tributária real pc", "=E{r}*S{r}/T{r}"),
        ("Transferências líquidas reais pc", "=O{r}*S{r}/T{r}"), ("Despesa real pc", "=M{r}*S{r}/T{r}")]
for mun in ["Manaus", "Parintins"]:
    ws = wb.create_sheet(mun)
    cab(ws, ["Ano"] + [l for _, l in ENT] + [l for l, _ in CALC])
    d = f[f.municipio == mun].sort_values("ano")
    for i, (_, r) in enumerate(d.iterrows(), 2):
        ws.cell(row=i, column=1, value=int(r.ano))
        if pd.isna(r.rec_corrente):
            ws.cell(row=i, column=2, value="sem DCA no Siconfi (ausente)")
            continue
        for j, (k, _) in enumerate(ENT, 2):
            v = r[k]
            c = ws.cell(row=i, column=j, value=None if pd.isna(v) else float(v)); c.fill = IN; c.number_format = "#,##0.00"
        for j, (_, fml) in enumerate(CALC, 2 + len(ENT)):
            c = ws.cell(row=i, column=j, value=fml.format(r=i, mun=mun)); c.number_format = "#,##0.00"

# ---- Conferência com Python
ws = wb.create_sheet("Conferencia")
cab(ws, ["Município", "Ano", "Dependência % (planilha)", "Dependência % (Python)", "Despesa real pc (planilha)", "Despesa real pc (Python)", "Diferença desp. pc"])
row = 2
for mun in ["Manaus", "Parintins"]:
    d = f[f.municipio == mun].sort_values("ano")
    for i, (_, r) in enumerate(d.iterrows(), 2):
        if pd.isna(r.rec_corrente):
            continue
        ws.cell(row=row, column=1, value=mun); ws.cell(row=row, column=2, value=int(r.ano))
        ws.cell(row=row, column=3, value=f"='{mun}'!P{i}"); ws.cell(row=row, column=4, value=float(100 * r.dep_transf))
        ws.cell(row=row, column=5, value=f"='{mun}'!X{i}"); ws.cell(row=row, column=6, value=float(r.desp_total_real_pc))
        ws.cell(row=row, column=7, value=f"=E{row}-F{row}")
        row += 1
saida = os.path.join(ROOT, "06_codigos_e_planilhas", "planilha_calculos_P2_231019943.xlsx")
wb.save(saida)
print("ok", saida)
